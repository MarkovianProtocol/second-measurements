#!/usr/bin/env python3
"""SM-025: pipeline operators' one-hour telephone notice to the National Response Center (49 CFR 191.5(a), 195.52(a)).
Inputs (DATA_DIR):
  PHMSA incident/accident bulk files (zips as downloaded): hazardous liquid, gas distribution, gas transmission/gathering.
    Each report carries the operator's "Local time of confirmed discovery" and the NRC report number.
  National Response Center FOIA call logs CY24.xlsx, CY25.xlsx, CY26.xlsx (sheet CALLS: SEQNOS = report number,
    DATE_TIME_RECEIVED in Eastern time).
Late = NRC receipt time, converted to the incident's local time zone, more than 65 minutes after confirmed discovery,
AND the operator's own form call time more than 60 minutes after it, AND the call's text cites no cost or volume
threshold (THRESH below). The plain one-hour count by the log alone is reported too (late_raw_over_60).
Counted: incidents dated 2024 onward whose report gives an NRC report number found in the call logs, where the NRC
record agrees with the PHMSA report on state and its incident date is within 2 days of the PHMSA incident time or of
the confirmed discovery. A number that belongs to a 48-hour
update is followed to the initial call it names. Reports marked
"NRC notification not required" are not counted; "required but not made" are counted separately.
python3 analyze.py DATA_DIR OUT.json [rows.csv]   (needs openpyxl)"""
import collections, csv, datetime as dt, io, json, os, re, statistics as st, sys, zipfile
from zoneinfo import ZoneInfo
import openpyxl
D, OUT = sys.argv[1], sys.argv[2]; ROWS = sys.argv[3] if len(sys.argv) > 3 else None
csv.field_size_limit(10 ** 9)
FILES = {"hazardous liquid": "accident_hazardous_liquid_jan2010_present", "gas distribution": "incident_gas_distribution_jan2010_present",
         "gas transmission and gathering": "incident_gas_transmission_gathering_jan2010_present"}
ZONES = {"EASTERN": "America/New_York", "CENTRAL": "America/Chicago", "MOUNTAIN": "America/Denver", "PACIFIC": "America/Los_Angeles",
         "ALASKA": "America/Anchorage", "HAWAII-ALEUTIAN": "Pacific/Honolulu"}
EAST = ZoneInfo("America/New_York")
BULLETIN = dt.date(2026, 2, 20)            # PHMSA ADB-2026-04 published
def p(s):
    for f in ("%m/%d/%Y %H:%M", "%m/%d/%Y %I:%M:%S %p"):
        try: return dt.datetime.strptime((s or "").strip(), f)
        except ValueError: pass
calls, inc = {}, {}
for y in ("24", "25", "26"):
    wb = openpyxl.load_workbook(os.path.join(D, f"CY{y}.xlsx"), read_only=True)
    it = wb["CALLS"].iter_rows(values_only=True); next(it)
    for r in it:
        if r[1]: calls[str(r[0])] = r[1]
    it = wb["INCIDENT_COMMONS"].iter_rows(values_only=True); h = next(it)
    ci, cs, cd = h.index("INCIDENT_DATE_TIME"), h.index("LOCATION_STATE"), h.index("DESCRIPTION_OF_INCIDENT")
    for r in it: inc[str(r[0])] = (r[ci], r[cs], r[cd] or "")
UPD = re.compile(r"(?:UPDATE|CORRECTION)\s+TO\s+(?:THE\s+)?(?:PREVIOUS\s+)?NRC\s+REPORT\s*(?:#|NO\.?|NUMBER)?\s*:?\s*#?\s*(\d{6,8})", re.I)
# Calls whose own text says the operator called because a cost or volume estimate crossed the reporting threshold, or that
# reportability was established after the release was found: for these the form's "confirmed discovery" may predate the
# moment the event became reportable, so they are shown separately, not counted as late in the main figure.
THRESH = re.compile(r"\$\s?\d|\b(?:50|122|129|139),?000\b|DAMAGE (?:AMOUNT|THRESHOLD|ESTIMATE)|EXCEED|THRESHOLD|CRITERIA|"
                    r"(?:BECAME|BECOME|DEEMED|DETERMINED|CONSIDERED|NOW|WAS NOT) (?:TO BE )?REPORTABLE|ESTIMATE[SD]? (?:THAT|THE)? ?(?:COST|DAMAGE|REPAIR)|"
                    r"COST (?:OF|TO) (?:THE )?(?:REPAIR|CLEAN)|REPAIR COST|RECALCULAT|RECHECK|SOIL SAMPLE|SURPASS|DOLLAR AMOUNT|CALLING BASED ON|"
                    r"REPORTABLE QUANTITY|\bRQ\b|IN EXCESS OF 3|3,000,000|THREE MILLION|INFORMED OF", re.I)
# The form's number is labelled an update but describes a new release (a second pipeline struck): keep the form's call.
NOFOLLOW = {("20260003", "HOLLY")}
# Reports where the NRC log shows an earlier call for the same incident that the form does not name (review gate, 24 Oct)
AMBIGUOUS = {("20240070", "TRUNKLINE"): "an earlier call from the same compressor station (NRC 1397734), about 48 minutes after discovery, is not on the form",
             ("20240188", "INLAND"): "an earlier call (NRC 1404683) reports diesel on the ground at the next address on the same road, about two hours before the counted call",
             ("20260093", "TTLENRGES"): "the counted call says it is related to NRC report 1456729, made 5 hours earlier, possibly by another company"}
def nrc_date(v):
    if isinstance(v, dt.datetime): return v.date()
    d = p(str(v or "")) or (lambda s: dt.datetime.strptime(s, "%m/%d/%Y") if re.fullmatch(r"\d+/\d+/\d{4}", s) else None)(str(v or "").strip())
    return d.date() if d else None
rows, notmade, unmatched, nodisc, total_reports, followed = [], [], [], 0, collections.Counter(), [0]
for system, name in FILES.items():
    z = zipfile.ZipFile(os.path.join(D, name + ".zip"))
    for r in csv.DictReader(io.TextIOWrapper(z.open(name + ".txt"), encoding="latin-1"), delimiter="\t"):
        if r["IYEAR"] not in ("2024", "2025", "2026"): continue
        total_reports[system] += 1
        num = (r["NRC_RPT_NUM"] or "").strip().upper()
        if "NOT REQUIRED" in num: continue
        if "REQUIRED BUT NOT MADE" in num: notmade.append({"system": system, "report": r["REPORT_NUMBER"], "operator": r["NAME"], "local": r["LOCAL_DATETIME"]}); continue
        m = re.search(r"\d{6,8}", num); disc = p(r["CONFIRMED_DISCOVERY_DATETIME"])
        if not m or not disc: nodisc += 1; continue
        state = r.get("ONSHORE_STATE_ABBREVIATION") or r.get("LOCATION_STATE_ABBREVIATION") or ""
        zname = "America/Phoenix" if state == "AZ" and r["TIME_ZONE"] == "MOUNTAIN" else ZONES.get(r["TIME_ZONE"])
        seq = m.group(0)
        if seq not in calls or not zname: unmatched.append({"system": system, "report": r["REPORT_NUMBER"], "nrc": seq, "why": "number not in NRC call logs"}); continue
        hops = 0
        while hops < 5 and (r["REPORT_NUMBER"], r["NAME"][:5]) not in NOFOLLOW:                                        # the form gives an update's or correction's number: follow it to the first call
            u = UPD.search(inc.get(seq, (None, None, ""))[2])
            if not (u and u.group(1) in calls and u.group(1) != seq): break
            seq = u.group(1); hops += 1
        if hops: followed[0] += 1
        amb = [v for (n, o), v in AMBIGUOUS.items() if n == r["REPORT_NUMBER"] and o in r["NAME"].upper()]
        if amb: unmatched.append({"system": system, "report": r["REPORT_NUMBER"], "operator": r["NAME"], "nrc": seq, "why": "ambiguous: " + amb[0]}); continue
        idate, istate, _ = inc.get(seq, (None, None, ""))
        local = p(r["LOCAL_DATETIME"]); nd = nrc_date(idate)
        why = ("state differs" if istate and state and istate != state else
               "NRC incident date more than 2 days from both the incident and its discovery" if nd and local and abs((nd - local.date()).days) > 2 and abs((nd - disc.date()).days) > 2 else
               "confirmed discovery more than a day before the incident time" if local and (local - disc).total_seconds() > 86400 else None)
        if why: unmatched.append({"system": system, "report": r["REPORT_NUMBER"], "operator": r["NAME"], "nrc": seq, "why": why}); continue
        tz = ZoneInfo(zname)
        recv_local = calls[seq].replace(tzinfo=EAST).astimezone(tz).replace(tzinfo=None)
        form = p(r["NRC_RPT_DATETIME"])
        fatal = int(float(r.get("FATAL") or 0)); injure = int(float(r.get("INJURE") or 0))
        rows.append({"system": system, "report": r["REPORT_NUMBER"], "operator_id": r["OPERATOR_ID"], "operator": r["NAME"], "state": state,
                     "time_zone": r["TIME_ZONE"], "confirmed_discovery_local": disc.isoformat(), "nrc_report": seq, "nrc_number_on_form": m.group(0),
                     "nrc_received_eastern": calls[seq].isoformat(), "nrc_received_local": recv_local.isoformat(),
                     "form_call_local": form.isoformat() if form else "", "minutes": round((recv_local - disc).total_seconds() / 60, 1),
                     "form_minus_log_min": round((form - recv_local).total_seconds() / 60, 1) if form else None,
                     "form_minutes": round((form - disc).total_seconds() / 60, 1) if form else None,
                     "threshold_language": bool(THRESH.search(inc.get(seq, (None, None, ""))[2])),
                     "nrc_description": re.sub(r"\s+", " ", inc.get(seq, (None, None, ""))[2])[:500],
                     "fatal": fatal, "injured": injure, "fire": r.get("IGNITE_IND") == "YES", "explosion": r.get("EXPLODE_IND") == "YES",
                     "cause": r.get("CAUSE", ""),
                     # 195.52(a)(1)-(4) on the final form values; None for gas. (a)(5), "in the judgment of the operator was
                     # significant", can't be read from the form, so a liquid call meeting none of (1)-(4) may still be required.
                     # Reportable on its face from the final form: death, injury, fire, explosion, 3 million cubic feet of gas
                     # (191.3) or water pollution (195.52(a)(4)); the rest were reportable on cost or the operator's judgment.
                     "reportable_on_its_face": bool(fatal or injure or r.get("IGNITE_IND") == "YES" or r.get("EXPLODE_IND") == "YES"
                                                    or (system != "hazardous liquid" and float(r.get("UNINTENTIONAL_RELEASE") or 0) >= 3000)
                                                    or (system == "hazardous liquid" and r.get("WATER_CONTAM_IND") == "YES")),
                     "liquid_criteria_1_to_4": (fatal > 0 or injure > 0 or r.get("IGNITE_IND") == "YES" or r.get("EXPLODE_IND") == "YES"
                                                or sum(float(r.get(c) or 0) for c in ("EST_COST_OPER_PAID", "EST_COST_GAS_RELEASED", "EST_COST_PROP_DAMAGE",
                                                       "EST_COST_EMERGENCY", "EST_COST_ENVIRONMENTAL", "EST_COST_OTHER")) > 50000
                                                or r.get("WATER_CONTAM_IND") == "YES") if system == "hazardous liquid" else None})
q = lambda xs, f: sorted(xs)[min(len(xs) - 1, int(f * len(xs)))]
# Main measure: late by the Center's log by more than 65 minutes (a 5-minute allowance for answer and hold time) AND late by
# the call time on the operator's own form, AND the call's text does not cite a cost or volume threshold.
for r in rows:
    r["late_raw"] = r["minutes"] > 60
    r["late_both"] = r["minutes"] > 65 and (r["form_minutes"] is None or r["form_minutes"] > 60)
    r["late"] = r["late_both"] and not r["threshold_language"]
    r["late_facial"] = r["late"] and r["reportable_on_its_face"]
def summ(rs):
    m = [r["minutes"] for r in rs]; late = [r["minutes"] for r in rs if r["late"]]
    if not m: return {"n": 0}
    return {"n": len(m), "late": len(late), "late_pct": round(100 * len(late) / len(m), 1), "over_2h": sum(x > 120 for x in late),
            "over_24h": sum(x > 1440 for x in late), "late_raw_over_60": sum(r["late_raw"] for r in rs), "late_by_both_over_65": sum(r["late_both"] for r in rs),
            "late_by_both_threshold_language": sum(r["late_both"] and r["threshold_language"] for r in rs),
            "reportable_on_its_face": sum(r["reportable_on_its_face"] for r in rs), "late_on_its_face": sum(r["late_facial"] for r in rs), "raw_over_24h": sum(x > 1440 for x in m), "within_1h": sum(0 <= x <= 60 for x in m), "before_discovery_15min": sum(x < -15 for x in m),
            "median_min": round(st.median(m), 1), "median_late_min": round(st.median(late), 1) if late else None,
            "p90_min": round(q(m, .9), 1)}
by_op = collections.defaultdict(list)
for r in rows: by_op[r["operator"]].append(r)
ops = sorted(({"operator": o, **summ(rs)} for o, rs in by_op.items() if len(rs) >= 5), key=lambda x: (-x["late"], -x["late_pct"]))
mm = [r["minutes"] for r in rows if 0 < r["minutes"] <= 180]
buckets = collections.Counter(min(int(x // 10) * 10, 170) for x in mm)
fl = [r["form_minus_log_min"] for r in rows if r["form_minus_log_min"] is not None]
serious = [r for r in rows if r["fatal"] or r["injured"] or r["fire"] or r["explosion"]]
out = {
    "window": "incidents dated 2024-01-01 onward (IYEAR 2024-2026) in PHMSA files of 14 Oct 2026; NRC logs CY2024-CY2026",
    "reports_in_window": dict(total_reports), "required_but_not_made": notmade, "no_number_or_no_discovery_time": nodisc,
    "not_counted": dict(collections.Counter(u["why"] for u in unmatched)), "not_counted_rows": unmatched, "update_numbers_followed_to_initial_call": followed[0],
    "all": summ(rows), "by_system": {s: summ([r for r in rows if r["system"] == s]) for s in FILES},
    "by_year": {y: summ([r for r in rows if r["confirmed_discovery_local"][:4] == y]) for y in ("2024", "2025", "2026")},
    "bulletin_gas_only": {"2024": summ([r for r in rows if r["system"] != "hazardous liquid" and r["confirmed_discovery_local"][:4] == "2024"]),
                          "2025": summ([r for r in rows if r["system"] != "hazardous liquid" and r["confirmed_discovery_local"][:4] == "2025"]),
                          "2026_before_2026-02-20": summ([r for r in rows if r["system"] != "hazardous liquid" and r["confirmed_discovery_local"][:4] == "2026" and dt.date.fromisoformat(r["confirmed_discovery_local"][:10]) < BULLETIN]),
                          "from_2026-02-20": summ([r for r in rows if r["system"] != "hazardous liquid" and dt.date.fromisoformat(r["confirmed_discovery_local"][:10]) >= BULLETIN])},
    "liquid_meeting_none_of_195_52_a1_to_a4": summ([r for r in rows if r["liquid_criteria_1_to_4"] is False]),
    "all_without_those": summ([r for r in rows if r["liquid_criteria_1_to_4"] is not False]),
    "late_by_form_alone_over_60": sum(1 for r in rows if r["form_minutes"] is not None and r["form_minutes"] > 60),
    "late_by_log_not_form": sum(1 for r in rows if r["minutes"] > 60 and r["form_minutes"] is not None and r["form_minutes"] <= 60),
    "late_by_form_not_log": sum(1 for r in rows if r["minutes"] <= 60 and r["form_minutes"] is not None and r["form_minutes"] > 60),
    "form_a_few_minutes_before_log": sum(1 for r in rows if r["form_minus_log_min"] is not None and -10 <= r["form_minus_log_min"] < 0),
    "form_a_few_minutes_after_log": sum(1 for r in rows if r["form_minus_log_min"] is not None and 0 < r["form_minus_log_min"] <= 10),
    "grace": {str(g): sum(r["minutes"] > 60 + g for r in rows) for g in (0, 5, 10, 15)},
    "serious_fatality_injury_fire_explosion": summ(serious), "fatal_or_injured": summ([r for r in rows if r["fatal"] or r["injured"]]),
    "operators_5plus": ops, "minutes_buckets_0_180": {f"{k}-{k + 10}": buckets[k] for k in sorted(buckets)},
    "form_vs_log": {"n": len(fl), "within_10min": sum(abs(x) <= 10 for x in fl), "off_by_about_an_hour_or_more": sum(abs(x) >= 50 for x in fl),
                    "median_abs_min": round(st.median(abs(x) for x in fl), 1)},
}
json.dump(out, open(OUT, "w"), indent=1)
if ROWS:
    with open(ROWS, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(sorted(rows, key=lambda r: (r["system"], r["confirmed_discovery_local"])))
a = out["all"]; print(f"facial {a['late_on_its_face']}/{a['reportable_on_its_face']}; {a['n']} calls: {a['late']} late ({a['late_pct']}%) [raw >60 by log {a['late_raw_over_60']}, late by both {a['late_by_both_over_65']}, of which threshold {a['late_by_both_threshold_language']}], over 2 h {a['over_2h']}, over 24 h {a['over_24h']}")
